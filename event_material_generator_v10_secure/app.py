import streamlit as st
import pandas as pd
from datetime import date
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from pathlib import Path
import sqlite3, json, hashlib, secrets, hmac, os, time
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from cryptography.fernet import Fernet, InvalidToken

st.set_page_config(page_title="Event Material Generator", page_icon="📋", layout="wide", initial_sidebar_state="auto")
ROOT = Path(__file__).parent
LETTERHEAD = ROOT / "assets_letterhead.docx"

# ---------------- Theme ----------------
st.markdown("""
<style>
:root{--navy:#06152d;--navy2:#0a2342;--blue:#1677b7;--purple:#5b35b5;--green:#087f5b;--bg:#eef2f6;--border:#d7e0e8;--text:#213247;--muted:#64748b}
html,body,[class*="css"]{font-family:Arial,sans-serif}
.stApp{background:var(--bg);color:var(--text);overflow-x:hidden}
@media (max-width: 768px){.block-container{padding:.65rem .75rem 4rem}.app-header{margin:-.65rem -.75rem .8rem;padding:10px 12px;display:block}.brand{font-size:18px}.brand-line{width:180px}.app-header>div:last-child{font-size:11px;margin-top:5px}.page-title{font-size:23px}.page-subtitle{font-size:13px}.section{padding:13px 12px;margin-bottom:12px}.section-title{font-size:18px}.card-value{font-size:18px}.stDataFrame,.stDataEditor{overflow-x:auto}.stButton>button,.stDownloadButton>button{min-height:42px;width:100%;white-space:normal}.stTextInput input,.stTextArea textarea,.stSelectbox div[data-baseweb='select']{font-size:16px!important}}
.block-container{padding-top:1rem;padding-bottom:2rem;max-width:1500px}
[data-testid="stSidebar"]{background:linear-gradient(180deg,var(--navy) 0%,#081d37 100%)}
[data-testid="stSidebar"]{color:#fff}
[data-testid="stSidebar"] .stRadio label{padding:9px 12px;border-radius:9px;background:#f4f7fa!important;color:#16324d!important;margin:5px 0;font-weight:700;border:1px solid #d9e2eb}
[data-testid="stSidebar"] .stRadio label *{color:#16324d!important}
[data-testid="stSidebar"] .stRadio label:hover{background:#ffffff!important;border-color:#a9c2d8}
[data-testid="stSidebar"] .stRadio [data-baseweb="radio"]{background:transparent!important}
[data-testid="stSidebar"] .stButton>button{background:#f5f7fa!important;color:#17324d!important;border:1px solid #d8e0e8!important;font-weight:700}
[data-testid="stSidebar"] .stButton>button:hover{background:#fff!important;color:#0b4f7a!important}
[data-testid="stSidebar"] .step-pill{background:#173553;color:#fff!important;border:1px solid #2d4b68}
.app-header{background:#fff;border-bottom:1px solid #dce3ea;display:flex;align-items:center;justify-content:space-between;padding:10px 20px;margin:-1rem -1rem 1rem -1rem}
.brand{font-weight:800;font-size:22px;color:#18293e;letter-spacing:.2px}.brand-line{height:4px;background:#e4a21b;width:440px;margin-top:4px}
.page-title{font-size:30px;font-weight:750;color:#24364b;margin:8px 0 3px}.page-subtitle{color:var(--muted);margin-bottom:18px}
.stButton>button{background:#173b5f!important;color:#ffffff!important;border:1px solid #173b5f!important;font-weight:700;border-radius:7px;min-height:40px}.stButton>button:hover{background:#0f2f4d!important;color:#ffffff!important}.stDownloadButton>button{background:#ffffff!important;color:#173b5f!important;border:1px solid #b8c7d6!important;font-weight:700;min-height:40px;border-radius:7px}.stDownloadButton>button:hover{background:#eef4f8!important;color:#0f2f4d!important}.section{background:#fff;border:1px solid var(--border);border-radius:8px;padding:18px 20px;margin-bottom:16px;box-shadow:0 1px 3px rgba(20,35,50,.05)}
.section-title{font-size:21px;font-weight:750;color:var(--purple);border-bottom:3px solid var(--green);padding-bottom:9px;margin-bottom:16px}
.req{color:#c82020;font-weight:800}.optional{color:#64748b;font-size:12px}
.card{background:#fff;border:1px solid var(--border);border-radius:8px;padding:15px;box-shadow:0 1px 3px rgba(20,35,50,.05)}
.card-title{font-size:13px;color:#64748b;text-transform:uppercase;letter-spacing:.5px}.card-value{font-size:22px;font-weight:800;color:#1c3147;margin-top:4px;overflow-wrap:anywhere}
.info-box{background:#edf8f4;border-left:4px solid var(--green);padding:11px 14px;border-radius:5px;color:#174e40;margin:8px 0 16px}.warning-box{background:#fff7e6;border-left:4px solid #e2a31a;padding:11px 14px;border-radius:5px;color:#6d5010}
.small-muted{font-size:12px;color:#718096}.site-footer{text-align:center;color:#64748b;font-size:12px;padding:22px 10px 8px;margin-top:28px;border-top:1px solid #d7e0e8}.mobile-stack{display:block}.history-row{background:#fff;border:1px solid var(--border);border-radius:8px;padding:14px;margin:8px 0}.history-title{font-weight:750;color:#1c3147;font-size:16px}.history-meta{color:#64748b;font-size:12px;margin-top:4px}.step-pill{display:inline-block;padding:5px 9px;border-radius:14px;background:#edf2f7;color:#334e68;font-size:11px;margin:2px}
</style>
""", unsafe_allow_html=True)

# ---------------- Persistent Database & Security ----------------
# IMPORTANT: secrets and persistent data are intentionally kept outside source control.
# For production/public deployment set APP_ENCRYPTION_KEY and ADMIN_PASSWORD in the
# hosting provider's secret manager (e.g. Streamlit secrets/environment variables).
DATA_DIR = Path(os.environ.get("EVENT_DATA_DIR", str(ROOT / ".private_data")))
DATA_DIR.mkdir(parents=True, exist_ok=True)
try:
    DATA_DIR.chmod(0o700)
except OSError:
    pass
DB_PATH = DATA_DIR / "event_material_generator.db"
KEY_PATH = DATA_DIR / "app_encryption.key"


def _secret(name, default=None):
    value = os.environ.get(name)
    if value:
        return value
    try:
        value = st.secrets.get(name)
        if value:
            return str(value)
    except Exception:
        pass
    return default


def get_fernet():
    key = _secret("APP_ENCRYPTION_KEY")
    if key:
        try:
            return Fernet(key.encode())
        except Exception as exc:
            raise RuntimeError("APP_ENCRYPTION_KEY is invalid. Generate a Fernet key and store it in your deployment secrets.") from exc
    # Local development convenience: generate a key in an ignored, permission-restricted directory.
    # Production deployments should always provide APP_ENCRYPTION_KEY via secrets.
    if KEY_PATH.exists():
        key = KEY_PATH.read_text(encoding="utf-8").strip()
    else:
        key = Fernet.generate_key().decode()
        KEY_PATH.write_text(key, encoding="utf-8")
        try:
            KEY_PATH.chmod(0o600)
        except OSError:
            pass
    return Fernet(key.encode())

FERNET = get_fernet()


def db_conn():
    con = sqlite3.connect(DB_PATH, timeout=10)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA foreign_keys=ON")
    return con


def hash_password(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 310000).hex()
    return f"pbkdf2_sha256$310000${salt}${digest}"


def verify_password(password, encoded):
    try:
        _, iters, salt, digest = encoded.split("$")
        calc = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), int(iters)).hex()
        return hmac.compare_digest(calc, digest)
    except Exception:
        return False


def init_db():
    con = db_conn()
    cur = con.cursor()
    cur.execute("CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, email TEXT UNIQUE, password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'user', active INTEGER NOT NULL DEFAULT 1, must_change_password INTEGER NOT NULL DEFAULT 0)")
    cur.execute("CREATE TABLE IF NOT EXISTS app_state (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
    cur.execute("CREATE TABLE IF NOT EXISTS event_history (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, start_date TEXT, end_date TEXT, owner_username TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL, payload TEXT NOT NULL)")
    # Safe migrations for databases created by earlier versions.
    cols={r[1] for r in cur.execute("PRAGMA table_info(users)").fetchall()}
    if 'email' not in cols: cur.execute("ALTER TABLE users ADD COLUMN email TEXT")
    if 'must_change_password' not in cols: cur.execute("ALTER TABLE users ADD COLUMN must_change_password INTEGER NOT NULL DEFAULT 0")
    # Existing accounts from pre-registration versions are retained. The legacy
    # administrator is marked for one-time secret bootstrap when needed.
    admin_username=_secret("ADMIN_USERNAME", "admin").strip() or "admin"
    row=cur.execute("SELECT * FROM users WHERE username=?",(admin_username,)).fetchone()
    configured_password=_secret("ADMIN_PASSWORD")
    configured_email=normalize_email(_secret("ADMIN_EMAIL")) if _secret("ADMIN_EMAIL") else ""
    if not row:
        bootstrap_path = DATA_DIR / "bootstrap_admin_password.txt"
        admin_password=configured_password
        if not admin_password:
            admin_password=secrets.token_urlsafe(18)
            bootstrap_path.write_text("Initial administrator password (change immediately): "+admin_password+"\n",encoding="utf-8")
            try: bootstrap_path.chmod(0o600)
            except OSError: pass
        cur.execute("INSERT INTO users(username,email,password_hash,role,active,must_change_password) VALUES(?,?,?,?,1,?)",(admin_username,configured_email or None,hash_password(admin_password),"admin",0))
    elif row['role']=='admin' and configured_email and not row['email']:
        cur.execute("UPDATE users SET email=? WHERE username=?",(configured_email,admin_username))
        if configured_password and int(row['must_change_password'] or 0)==1:
            cur.execute("UPDATE users SET password_hash=?,must_change_password=0 WHERE username=?",(hash_password(configured_password),admin_username))
    elif row['role']=='admin' and configured_password and ('must_change_password' in cols) and int(row['must_change_password'] or 0)==1:
        cur.execute("UPDATE users SET password_hash=?,must_change_password=0 WHERE username=?",(hash_password(configured_password),admin_username))
    elif row['role']=='admin' and 'must_change_password' not in cols:
        cur.execute("UPDATE users SET must_change_password=1 WHERE username=?",(admin_username,))
        if configured_password:
            cur.execute("UPDATE users SET password_hash=?,must_change_password=0 WHERE username=?",(hash_password(configured_password),admin_username))
    con.commit(); con.close()

def normalize_email(email):
    return email.strip().lower()

def valid_email(email):
    """Registration/admin email must be an NFSU institutional address."""
    import re
    return bool(re.fullmatch(r"[^\s@]+@nfsu\.ac\.in", normalize_email(email)))

def create_user(username,email,password,role='user',must_change=False):
    username=username.strip()
    email=normalize_email(email)
    con=db_conn()
    try:
        con.execute("INSERT INTO users(username,email,password_hash,role,active,must_change_password) VALUES(?,?,?,?,1,?)",(username,email,hash_password(password),role,1 if must_change else 0))
        con.commit()
        return True,None
    except sqlite3.IntegrityError as ex:
        return False,str(ex)
    finally: con.close()

def save_event_history(event, owner_username, event_id=None):
    payload=serialize_event_obj(event)
    now=time.strftime('%Y-%m-%d %H:%M:%S')
    con=db_conn()
    if event_id:
        con.execute("UPDATE event_history SET title=?,start_date=?,end_date=?,owner_username=?,updated_at=?,payload=? WHERE id=?",(event.get('title','Untitled Event'),str(event.get('start_date','')),str(event.get('end_date','')),owner_username,now,payload,int(event_id)))
    else:
        cur=con.execute("INSERT INTO event_history(title,start_date,end_date,owner_username,created_at,updated_at,payload) VALUES(?,?,?,?,?,?,?)",(event.get('title','Untitled Event'),str(event.get('start_date','')),str(event.get('end_date','')),owner_username,now,now,payload))
        event_id=cur.lastrowid
    con.commit(); con.close(); return int(event_id)

def list_event_history(owner_username=None):
    con=db_conn()
    if owner_username:
        rows=con.execute("SELECT id,title,start_date,end_date,owner_username,created_at,updated_at FROM event_history WHERE owner_username=? ORDER BY COALESCE(start_date,'') DESC,id DESC",(owner_username,)).fetchall()
    else:
        rows=con.execute("SELECT id,title,start_date,end_date,owner_username,created_at,updated_at FROM event_history ORDER BY COALESCE(start_date,'') DESC,id DESC").fetchall()
    con.close(); return [dict(r) for r in rows]

def get_event_history(event_id):
    con=db_conn(); row=con.execute("SELECT * FROM event_history WHERE id=?",(int(event_id),)).fetchone(); con.close()
    return dict(row) if row else None

def delete_event_history(event_id):
    con=db_conn(); con.execute("DELETE FROM event_history WHERE id=?",(int(event_id),)); con.commit(); con.close()


def encrypt_text(value):
    return FERNET.encrypt(value.encode("utf-8")).decode("ascii")


def decrypt_text(value):
    try:
        return FERNET.decrypt(value.encode("ascii")).decode("utf-8")
    except (InvalidToken, ValueError, TypeError):
        # Migration support for older local databases that stored JSON in plaintext.
        return value


def db_get(key, default=None):
    con = db_conn()
    row = con.execute("SELECT value FROM app_state WHERE key=?", (key,)).fetchone()
    con.close()
    if not row:
        return default
    raw = decrypt_text(row["value"])
    try:
        value = json.loads(raw)
    except Exception:
        return default
    # If this was legacy plaintext, transparently migrate it to encrypted storage.
    if raw == row["value"]:
        try:
            db_set(key, value)
        except Exception:
            pass
    return value


def db_set(key, value):
    raw = json.dumps(value, default=str, ensure_ascii=False)
    encrypted = encrypt_text(raw)
    con = db_conn()
    con.execute(
        "INSERT INTO app_state(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, encrypted),
    )
    con.commit()
    con.close()
    try:
        DB_PATH.chmod(0o600)
    except OSError:
        pass


def db_users():
    con = db_conn()
    rows = con.execute("SELECT id,username,email,role,active,must_change_password FROM users ORDER BY username").fetchall()
    con.close()
    return [dict(r) for r in rows]


def authenticate(username, password):
    con = db_conn()
    identifier=username.strip()
    row = con.execute("SELECT * FROM users WHERE active=1 AND (username=? OR lower(COALESCE(email,''))=?)", (identifier, normalize_email(identifier))).fetchone()
    con.close()
    return dict(row) if row and verify_password(password, row["password_hash"]) else None

def update_user(user_id, email, role, active):
    con=db_conn(); con.execute("UPDATE users SET email=?,role=?,active=? WHERE id=?",(normalize_email(email) if email else None,role,1 if active else 0,int(user_id))); con.commit(); con.close()

def reset_user_password(user_id, new_password):
    con=db_conn(); con.execute("UPDATE users SET password_hash=?,must_change_password=1 WHERE id=?",(hash_password(new_password),int(user_id))); con.commit(); con.close()

def change_password(username,new_password):
    con=db_conn(); con.execute("UPDATE users SET password_hash=?,must_change_password=0 WHERE username=?",(hash_password(new_password),username)); con.commit(); con.close()

# Simple per-process login throttling. Production deployments should also use a
# reverse proxy/WAF rate limit when the app is internet-facing.
_LOGIN_FAILURES = {}

def login_allowed(username):
    rec = _LOGIN_FAILURES.get(username.strip().lower(), {"count": 0, "until": 0})
    return time.time() >= rec["until"]


def record_login_failure(username):
    key = username.strip().lower()
    rec = _LOGIN_FAILURES.get(key, {"count": 0, "until": 0})
    rec["count"] += 1
    if rec["count"] >= 5:
        rec["until"] = time.time() + min(300, 30 * (2 ** min(rec["count"] - 5, 3)))
    _LOGIN_FAILURES[key] = rec


def clear_login_failures(username):
    _LOGIN_FAILURES.pop(username.strip().lower(), None)

def serialize_event_obj(event):
    def conv(o):
        if isinstance(o,date): return o.isoformat()
        if isinstance(o,pd.DataFrame): return {"__dataframe__":True,"records":o.to_dict("records")}
        raise TypeError
    return json.dumps(event,default=conv)

def deserialize_event_obj(raw):
    obj=json.loads(raw)
    for k in ["start_date","end_date"]:
        if k in obj: obj[k]=date.fromisoformat(obj[k])
    for k in ["schedule_df","inauguration_df"]:
        v=obj.get(k)
        if isinstance(v,dict) and v.get("__dataframe__"): obj[k]=pd.DataFrame(v.get("records",[]))
        elif isinstance(v,list): obj[k]=pd.DataFrame(v)
    return obj

def save_current_persistent():
    if not st.session_state.get("authenticated"): return
    ev=st.session_state.get("event")
    if not ev:return
    ev={**ev,"schedule_df":st.session_state.get("schedule_df",pd.DataFrame()),"inauguration_df":st.session_state.get("inauguration_df",pd.DataFrame()),"experts":st.session_state.get("experts",[]),"budget":st.session_state.get("budget",[]),"checklist":st.session_state.get("checklist",{}),"chief_guest":st.session_state.get("chief_guest",{}),"coordinator_signature":st.session_state.get("coordinator_signature",""),"invitation_text":st.session_state.get("invitation_text","")}
    st.session_state.event=ev
    db_set("current_event_" + st.session_state.user.get("username",""),serialize_event_obj(ev)); db_set("faculty",st.session_state.get("faculty",[])); db_set("faculty_master",st.session_state.get("faculty_master",[])); db_set("event_heads",st.session_state.get("event_heads",[])); db_set("leadership",st.session_state.get("leadership",{})); db_set("institution",st.session_state.get("institution",{}))
    if st.session_state.get("current_event_id") and st.session_state.get("user"):
        save_event_history(ev, st.session_state.user.get("username"), st.session_state.current_event_id)

init_db()

# ---------------- State ----------------
def demo_event():
    schedule = pd.DataFrame({
        "Day":["1","1","1","2","2","2","3","3","3"],
        "Time":["09:30–10:00","10:00–11:30","11:45–13:00","09:30–11:00","11:15–12:45","14:00–15:30","09:30–11:00","11:15–12:45","14:00–15:30"],
        "Topic":["Registration & Inauguration","Introduction to Gait Pattern Analysis","Gait Cycle and Spatiotemporal Parameters","Clinical and Forensic Gait Assessment","Motion Capture and Data Acquisition","Practical: Gait Data Collection","AI-Assisted Gait Analysis","Practical: Feature Extraction and Interpretation","Case Study, Assessment and Valedictory"],
        "Expert Name":["Organising Team","Dr. Ananya Rao","Dr. Ananya Rao","Prof. Meera Nair","Dr. Rohan Kulkarni","Dr. Rohan Kulkarni","Prof. Meera Nair","Prof. Meera Nair","All Experts"]
    })
    inauguration = pd.DataFrame({"Time":["09:30","09:35","09:40","09:50","10:00"],"Activity":["Arrival of Guests","Welcome Address","Lighting of Lamp","Inaugural Address","Introduction to Workshop"],"Person":["Protocol Team","Coordinator","Chief Guest","Chief Guest","Coordinator"]})
    budget=[
        {"item":"Participation Fees","amount":0.0,"remarks":"No fee for invited participants"},
        {"item":"Honorarium to the Experts/Speakers","amount":15000.0,"remarks":"As per institutional norms"},
        {"item":"TA/DA & stay arrangements of the Experts/Speakers","amount":10000.0,"remarks":"As applicable"},
        {"item":"Boarding and Lodging/Hospitality to participants","amount":10000.0,"remarks":"Refreshments and working lunch"},
        {"item":"Contingency, Stationery etc.","amount":3000.0,"remarks":"Certificates, banner and stationery"},
        {"item":"Miscellaneous Expenses","amount":2000.0,"remarks":"Other incidental expenses"},
    ]
    checklist={m:[{"task":x,"status":"Not Started","responsible":"","due":"","remarks":""} for x in tasks] for m,tasks in CHECKLISTS.items()}
    return {
        "institution":{"name":"National Forensic Sciences University, Goa Campus","department":"School / Department","address":"Goa Campus, National Forensic Sciences University","email":"","phone":""},
        "programme_type":"Workshop","title":"Three-Day Workshop on Gait Pattern Analysis","coordinators":["Dr. Ranjit Kolkar"],"co_coordinators":["Dr. Ananya Rao"],"event_head":"Dean Academics",
        "start_date":date(2026,10,15),"end_date":date(2026,10,17),"mode":"Offline","venue":"NFSU Goa Campus","participants":40,
        "nature_list":["Theoretical","Practical","Workshop","Hands-on"],"audience":"Faculty, Researchers, Forensic Science Students, Law Enforcement and allied professionals",
        "infrastructure":"Forensic science laboratory facilities, computing systems, image/video analysis workstations, motion analysis setup and seminar hall facilities available at the campus.",
        "organizer_expertise":"The organizing team has expertise in forensic science, digital evidence, data analysis, artificial intelligence and applied research relevant to gait pattern assessment.",
        "introduction":"Gait pattern analysis is an interdisciplinary area with applications in forensic science, biometric identification, rehabilitation, security and investigative practice. Advances in motion analysis, computer vision and artificial intelligence have created new opportunities for objective extraction and interpretation of gait characteristics.",
        "justification":"There is a growing need to provide faculty members, researchers, students and professional stakeholders with structured exposure to the principles, practical methods and emerging computational approaches used in gait pattern analysis. A focused workshop will help bridge theoretical concepts with hands-on data acquisition, feature extraction and interpretation, while encouraging interdisciplinary collaboration and responsible application of analytical methods.",
        "objectives":"The workshop aims to introduce participants to the gait cycle and measurable gait parameters; demonstrate methods for gait data acquisition and analysis; provide practical exposure to feature extraction and interpretation; discuss applications of artificial intelligence and computer vision in gait analysis; and develop awareness of forensic, scientific and ethical considerations in the use of gait evidence.",
        "outcome":"Participants are expected to gain foundational and practical understanding of gait pattern analysis, develop familiarity with relevant analytical workflows and tools, and understand how gait-related features may be examined in research, forensic and professional contexts. The programme is also expected to promote collaborative research and further skill development in AI-assisted gait analysis.",
        "lodging":"Limited institutional accommodation may be arranged subject to availability; nearby hotels may be used where required.","sponsors":"Institutional support; external sponsorship may be explored if available.","platform":"","other_info":"Three-day programme with lectures, demonstrations, practical sessions and case discussion.",
        "schedule_df":schedule,"inauguration_df":inauguration,"experts":[
            {"name":"Dr. Ananya Rao","affiliation":"Forensic Science Research Centre","specialization":"Gait analysis and forensic biomechanics","topic":"Introduction to Gait Pattern Analysis","role":"Resource Person"},
            {"name":"Dr. Rohan Kulkarni","affiliation":"Institute of Biomedical Analytics","specialization":"Motion capture and data analysis","topic":"Motion Capture and Gait Data Acquisition","role":"Resource Person"},
            {"name":"Prof. Meera Nair","affiliation":"Centre for AI and Computer Vision","specialization":"AI-assisted image and movement analysis","topic":"AI-Assisted Gait Analysis","role":"Resource Person"}],
        "chief_guest":{"name":"Prof. Arun Mehta","designation":"Professor and Senior Expert","affiliation":"National Research Institute","address":"New Delhi, India","role":"Chief Guest"},
        "budget":budget,"checklist":checklist,"coordinator_signature":"Dr. Ranjit Kolkar\nCoordinator, NFSU Goa","invitation_text":"We are pleased to invite you to grace the Three-Day Workshop on Gait Pattern Analysis and share your valuable guidance with the participants."
    }

CHECKLISTS={
    "Pre-Event":["Finalize event title, date and venue","Confirm coordinator / co-coordinator","Prepare official proposal","Prepare Note for Approval","Finalize budget estimate","Identify experts / speakers","Send expert invitations","Confirm Chief Guest / special guests","Finalize course content and schedule","Prepare inauguration programme","Prepare participant invitation / circular","Create registration form / link","Prepare poster / banner","Arrange venue and seating","Arrange audio-visual equipment","Arrange refreshments / hospitality","Arrange accommodation / transport if required","Prepare attendance sheet","Prepare certificates / mementos","Brief volunteers and event team"],
    "In-Event":["Registration desk operational","Attendance recorded","Chief Guest / experts received","Inauguration conducted as scheduled","Sessions conducted as per programme","Timekeeping / session coordination","Photography completed","Videography / recording completed","Refreshments / hospitality managed","Certificates / mementos distributed","Feedback collected","Important documents / photographs backed up"],
    "Post-Event":["Finalize attendance","Collect pending bills / vouchers","Record actual expenses","Process honorarium / reimbursements","Send thank-you letters","Compile photographs / videos","Prepare event report","Analyze participant feedback","Prepare media / website / social media content","Archive event documents","Submit completion documentation"]}

def blank_event():
    return {
        "institution":{"name":"National Forensic Sciences University, Goa Campus","department":"","address":"","email":"","phone":""},
        "programme_type":"Workshop","title":"","coordinators":[],"co_coordinators":[],"event_head":"Dean Academics",
        "start_date":date.today(),"end_date":date.today(),"mode":"Offline","venue":"","participants":1,
        "nature_list":[],"audience":"","infrastructure":"","organizer_expertise":"","introduction":"","justification":"","objectives":"","outcome":"",
        "lodging":"","sponsors":"","platform":"","other_info":"","schedule_df":pd.DataFrame(columns=["Day","Time","Topic","Expert Name"]),
        "inauguration_df":pd.DataFrame(columns=["Time","Activity","Person"]),"experts":[],"chief_guest":{"name":"","designation":"","affiliation":"","address":"","role":"Chief Guest"},
        "budget":[],"checklist":{m:[{"task":x,"status":"Not Started","responsible":"","due":"","remarks":""} for x in items] for m,items in CHECKLISTS.items()},
        "coordinator_signature":"","invitation_text":""
}

def init_state():
    d=blank_event()
    user_obj=st.session_state.get("user") or {}
    username=user_obj.get("username") or "__anon__"
    saved=db_get("current_event_" + username)
    if saved:
        try:d=deserialize_event_obj(saved)
        except Exception:pass
    faculty_master=db_get("faculty_master") or [
        {"name":"Prof. (Dr.) Naveen Kumar Chaudhary","designation":"Professor"},{"name":"Dr. Narayan Pandurang Waghmare","designation":"Professor"},{"name":"Dr. Lokesh Chouhan","designation":"Associate Professor"},{"name":"Dr. Raj Kumar Jaiswal","designation":"Assistant Professor"},{"name":"Dr. Panem Charanarur","designation":"Assistant Professor"},{"name":"Dr. Jerin Mohan N D","designation":"Assistant Professor"},{"name":"Dr. Sneha Sagar","designation":"Assistant Professor"},{"name":"Dr. Suryakant A. Patil","designation":"Assistant Professor"},{"name":"Dr. Mrinmayee Kale","designation":"Assistant Professor"},{"name":"Dr. Sweta Nidhi","designation":"Assistant Professor"},{"name":"Dr. Rakhee Lohia","designation":"Assistant Professor"},{"name":"Dr. Ranjit Kolkar","designation":"Assistant Professor"},{"name":"Dr. T.K. Gundoor","designation":"Assistant Professor"},{"name":"Dr. Inder Bhan Singh","designation":"Assistant Professor"},{"name":"Dr. Manpreet Kaur","designation":"Assistant Professor"},{"name":"Dr. Jovi Jose Salvador D'silva","designation":"Assistant Professor"},{"name":"Dr. Pranitha Sanda","designation":"Assistant Professor"},{"name":"Dr. Seema Malhotra","designation":"Assistant Professor"},{"name":"Dr. Sujit Mahato","designation":"Assistant Professor"},{"name":"Dr. Sanay Naha","designation":"Assistant Professor"},{"name":"Mr. Dova Nani","designation":"Lecturer"},{"name":"Mr. Harsh Panchal","designation":"Lecturer"},{"name":"Mr. Rajesh Kumar Mitra","designation":"Lecturer"},{"name":"Mr. Rahul Kamble","designation":"Lecturer"},{"name":"Ms. Anouska Dutta","designation":"Lecturer"},{"name":"Mr. Abhinav Salgunan","designation":"Lecturer"}]
    faculty=[x["name"] for x in faculty_master]
    heads=db_get("event_heads") or ["Dean Academics","Campus Director","Head of Department"]
    leadership=db_get("leadership") or {"Chief Patron":"Padmashri Dr. J. M. Vyas","Chair":"Dr. Naveen Kumar Choudhary","Co-Chair":"Dr. Lokesh Chouhan"}
    defaults={"event":d,"schedule_mode":"Day-wise","schedule_df":d.get("schedule_df",pd.DataFrame()).copy(),"inauguration_df":d.get("inauguration_df",pd.DataFrame()).copy(),"experts":d.get("experts",[]).copy(),"budget":d.get("budget",[]).copy(),"checklist":d.get("checklist",{}).copy(),"chief_guest":d.get("chief_guest",{}).copy(),"coordinator_signature":d.get("coordinator_signature",""),"invitation_text":d.get("invitation_text",""),"institution":d.get("institution",{}).copy(),"faculty":faculty,"faculty_master":faculty_master,"event_heads":heads,"leadership":leadership,"generated":None,"authenticated":False,"user":None,"current_event_id":None}
    for k,v in defaults.items():
        if k not in st.session_state:st.session_state[k]=v
init_state()

# ---------------- Helpers ----------------
def money(n): return f"₹{float(n):,.0f}/-"
def amount_words(n):
    n=int(round(float(n)))
    if n==0:return "Zero Only"
    ones=['','One','Two','Three','Four','Five','Six','Seven','Eight','Nine','Ten','Eleven','Twelve','Thirteen','Fourteen','Fifteen','Sixteen','Seventeen','Eighteen','Nineteen']; tens=['','','Twenty','Thirty','Forty','Fifty','Sixty','Seventy','Eighty','Ninety']
    def two(x): return ones[x] if x<20 else tens[x//10]+((' '+ones[x%10]) if x%10 else '')
    def three(x): return two(x) if x<100 else ones[x//100]+' Hundred'+((' '+two(x%100)) if x%100 else '')
    parts=[]
    for val,name in [(n//10000000,'Crore'),((n//100000)%100,'Lakh'),((n//1000)%100,'Thousand'),((n//100)%10,'Hundred'),(n%100,'')]:
        if val: parts.append(three(val)+((' '+name) if name else ''))
    return ' '.join(parts)+' Only'
def duration_text(s,e): return s.strftime('%d %B %Y') if s==e else f"{s.strftime('%d %B %Y')} to {e.strftime('%d %B %Y')}"
def clean_schedule(df):
    if df is None:return pd.DataFrame()
    out=df.fillna('').astype(str)
    return out.loc[out.apply(lambda r:any(str(x).strip() for x in r),axis=1)].reset_index(drop=True)

def set_cell_margins(cell,top=70,start=70,bottom=70,end=70):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def add_shading(cell,fill='E6EEF5'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def new_doc():
    if LETTERHEAD.exists():
        doc=Document(str(LETTERHEAD))
        body=doc._element.body
        for child in list(body):
            if child.tag != qn('w:sectPr'): body.remove(child)
        # Letterhead sample uses 1-inch margins; retain them.
    else: doc=Document()
    sec=doc.sections[0]; sec.top_margin=Cm(2.54); sec.bottom_margin=Cm(2.54); sec.left_margin=Cm(2.54); sec.right_margin=Cm(2.54)
    doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(10.5)
    return doc

def add_para(doc,text='',bold=False,center=False,size=None,space_after=6):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(space_after)
    r=p.add_run(str(text)); r.bold=bold
    if size:r.font.size=Pt(size)
    r.font.name='Arial'; return p

def add_kv(doc,label,value):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(label); r.bold=True; r.font.name='Arial'; r.font.size=Pt(10.5); rr=p.add_run(str(value or '')); rr.font.name='Arial'; rr.font.size=Pt(10.5)

def build_table(doc,df,font_size=9):
    df=clean_schedule(df)
    if df.empty:return
    t=doc.add_table(rows=1,cols=len(df.columns)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(df.columns):
        c=t.rows[0].cells[i]; c.text=str(h); add_shading(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.bold=True; r.font.name='Arial'; r.font.size=Pt(font_size)
    for _,row in df.iterrows():
        cells=t.add_row().cells
        for i,v in enumerate(row.tolist()):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                for r in p.runs:r.font.name='Arial'; r.font.size=Pt(font_size)
    return t

def build_note(event):
    doc=new_doc(); add_para(doc,'Note',bold=True,center=True,size=14,space_after=10)
    for key in ['introduction','justification','objectives','outcome']:
        if event.get(key): add_para(doc,event[key],space_after=9)
    total=sum(float(x.get('amount',0) or 0) for x in event.get('budget',[]))
    inst=event.get('institution',{}).get('name','the Institute'); title=event.get('title','the proposed programme'); mode=event.get('mode','Offline'); dates=duration_text(event['start_date'],event['end_date']); venue=event.get('venue','the Institute')
    add_para(doc,f"Approval is kindly sought to conduct the above-mentioned programme at {venue} and to allocate the necessary budget of approximately Rupees: {total:,.0f}/- ({amount_words(total)}).",space_after=12)
    add_para(doc,'Date: ____________________',space_after=18); add_para(doc,'Signature of the Coordinator(s)')
    return doc

def build_proposal(event):
    doc=new_doc(); add_para(doc,'Format for submitting proposals for organizing Workshop/e-Workshop',bold=True,center=True,size=13,space_after=10)
    add_kv(doc,'Type of Proposed Programme: ',event.get('programme_type')); add_kv(doc,'1. Name of Coordinator(s)/Convener: ',', '.join(event.get('coordinators',[]))); add_kv(doc,'2. Name of the Co-Coordinator(s)/Convener: ',', '.join(event.get('co_coordinators',[]))); add_kv(doc,'3. Title of the proposed programme: ',event.get('title')); add_kv(doc,'4. Duration & dates of the programme: ',duration_text(event['start_date'],event['end_date'])); add_kv(doc,'5. Nature of the Programme: ',', '.join(event.get('nature_list',[]))); add_kv(doc,'Details of available Infrastructure/Laboratory facilities related to the programme: ',event.get('infrastructure')); add_kv(doc,'Expertise of the organizer(s): ',event.get('organizer_expertise')); add_kv(doc,'6. Objectives of the Programme: ',event.get('objectives')); add_kv(doc,'7. Brief justification of proposal: ',event.get('justification')); add_kv(doc,'8. Course content (date-wise tentative schedule): ',''); build_table(doc,event.get('schedule_df'))
    add_kv(doc,'9. Number of participants expected / Targeted Audience: ',f"{event.get('participants')} / {event.get('audience')}"); add_kv(doc,'Details of lodging facilities: ',event.get('lodging')); total=sum(float(x.get('amount',0) or 0) for x in event.get('budget',[])); add_kv(doc,'11. Financial assistance required: ',money(total)); add_kv(doc,'12. Agencies expected to sponsor/finance the event: ',event.get('sponsors')); add_kv(doc,'13. Details of Probable Experts: ','')
    experts=pd.DataFrame([{'S. No.':i,'Name of Expert':x.get('name',''),'Affiliation':x.get('affiliation',''),'Field of Interest / Specialization':x.get('specialization',''),'Proposed Topic':x.get('topic','')} for i,x in enumerate(event.get('experts',[]),1)]); build_table(doc,experts,8.5)
    add_kv(doc,'14. Budget Estimates: ',''); bdf=pd.DataFrame([{'Sr. No.':i,'Item Description':x.get('item',''),'Amount in Rs.':f"{float(x.get('amount',0) or 0):,.2f}",'Remarks (if any)':x.get('remarks','')} for i,x in enumerate(event.get('budget',[]),1)]); build_table(doc,bdf,8.5); add_para(doc,f"TOTAL: ₹{total:,.2f}",bold=True,space_after=8)
    add_kv(doc,'16. Platform to be used for online event (if online): ',event.get('platform')); add_kv(doc,'17. Other information, if any: ',event.get('other_info')); add_para(doc,'Date: ____________________',space_after=14); add_para(doc,'Signature of the Coordinator(s)        (Member)        (Member)',space_after=14); add_para(doc,'Dean (Academics)',space_after=14); add_para(doc,'Campus Director')
    return doc

def build_invitation(event):
    doc=new_doc(); add_para(doc,'INVITATION',bold=True,center=True,size=14,space_after=10); g=event.get('chief_guest',{}); add_para(doc,f"To\n{g.get('name','')}\n{g.get('designation','')}\n{g.get('affiliation','')}\n{g.get('address','')}",space_after=10); add_para(doc,f"Subject: Invitation to grace the {event.get('title','')}",bold=True,space_after=8); text=event.get('invitation_text') or f"We are pleased to invite you to the {event.get('title','')} organized by {event.get('institution',{}).get('name','the Institute')}. The programme is scheduled from {duration_text(event['start_date'],event['end_date'])} at {event.get('venue','')}."; add_para(doc,text,space_after=9); add_para(doc,f"We would be honoured by your gracious presence and request you to grace the programme as the {g.get('role','Chief Guest')}.",space_after=9); add_para(doc,'We look forward to your kind presence and valuable guidance.',space_after=14); add_para(doc,'Yours sincerely,'); add_para(doc,event.get('coordinator_signature',''))
    return doc

def build_checklist_doc(event,module):
    doc=new_doc(); add_para(doc,f'{module} Checklist',bold=True,center=True,size=14,space_after=8); add_para(doc,f"Event: {event.get('title','Event')}",center=True,space_after=4); add_para(doc,f"Date: {duration_text(event['start_date'],event['end_date'])}",center=True,space_after=12)
    items=event.get('checklist',{}).get(module,[]); df=pd.DataFrame([{'✓':'☐','Task':x.get('task',''),'Status':x.get('status','Not Started'),'Responsible':x.get('responsible',''),'Due Date':x.get('due',''),'Remarks':x.get('remarks','')} for x in items]); build_table(doc,df,8.5); return doc

def build_inauguration(event):
    doc=new_doc(); add_para(doc,'Inauguration Programme',bold=True,center=True,size=14,space_after=8); add_para(doc,event.get('title',''),center=True,space_after=4); add_para(doc,duration_text(event['start_date'],event['end_date']),center=True,space_after=12); build_table(doc,event.get('inauguration_df'),9); return doc

def doc_bytes(doc): b=BytesIO(); doc.save(b); return b.getvalue()

def build_package(event):
    files={
        '01_Official_Workshop_Proposal.docx':doc_bytes(build_proposal(event)),
        '02_Note_for_Approval.docx':doc_bytes(build_note(event)),
        '03_Programme_Schedule.docx':doc_bytes(build_schedule_doc(event)),
        '04_Chief_Guest_Invitation.docx':doc_bytes(build_invitation(event)),
        '05_Inauguration_Programme.docx':doc_bytes(build_inauguration(event)),
        '06_Checklist_Pre_Event.docx':doc_bytes(build_checklist_doc(event,'Pre-Event')),
        '07_Checklist_In-Event.docx':doc_bytes(build_checklist_doc(event,'In-Event')),
        '08_Checklist_Post_Event.docx':doc_bytes(build_checklist_doc(event,'Post-Event')),
    }
    with BytesIO() as m:
        with ZipFile(m,'w',ZIP_DEFLATED) as z:
            for name,data in files.items(): z.writestr(name,data)
        return m.getvalue(),files

def build_schedule_doc(event):
    doc=new_doc(); add_para(doc,'Programme Schedule',bold=True,center=True,size=14,space_after=8); add_para(doc,event.get('title',''),center=True,space_after=4); add_para(doc,duration_text(event['start_date'],event['end_date']),center=True,space_after=12); build_table(doc,event.get('schedule_df'),9); return doc

def serialize_event(event):
    import json
    def conv(o):
        if isinstance(o,date): return o.isoformat()
        if isinstance(o,pd.DataFrame): return o.to_dict('records')
        return o
    return json.dumps(event,default=conv,indent=2).encode()

def load_event_dict(raw):
    import json
    obj=json.loads(raw.decode('utf-8'))
    for k in ['start_date','end_date']:
        if k in obj: obj[k]=date.fromisoformat(obj[k])
    for k in ['schedule_df','inauguration_df']:
        if isinstance(obj.get(k),list): obj[k]=pd.DataFrame(obj[k])
    return obj

def save_event_to_state(ev):
    st.session_state.event=ev; st.session_state.institution=ev.get('institution',st.session_state.institution); st.session_state.schedule_df=ev.get('schedule_df',st.session_state.schedule_df); st.session_state.inauguration_df=ev.get('inauguration_df',st.session_state.inauguration_df); st.session_state.experts=ev.get('experts',[]); st.session_state.budget=ev.get('budget',[]); st.session_state.checklist=ev.get('checklist',st.session_state.checklist); st.session_state.chief_guest=ev.get('chief_guest',{}); st.session_state.coordinator_signature=ev.get('coordinator_signature',''); st.session_state.invitation_text=ev.get('invitation_text','')

# ---------------- Authentication ----------------
if not st.session_state.get("authenticated",False):
    st.markdown("<div style='max-width:460px;margin:50px auto;background:#fff;border:1px solid #d7e0e8;border-radius:12px;padding:28px;box-shadow:0 8px 30px rgba(20,35,50,.08)'><div style='font-size:25px;font-weight:800;color:#18293e'>EVENT MATERIAL GENERATOR</div><div style='color:#64748b;margin:6px 0 22px'>Institutional Event Planning & Documentation</div>",unsafe_allow_html=True)
    login_tab, register_tab = st.tabs(["Sign in","Register"])
    with login_tab:
        with st.form("login_form"):
            u=st.text_input("Username or Email",placeholder="Enter username or email")
            pw=st.text_input("Password",type="password",placeholder="Enter password")
            submitted=st.form_submit_button("Sign in",type="primary",use_container_width=True)
        if submitted:
            if not login_allowed(u):
                st.error("Too many failed attempts. Please wait a few minutes before trying again.")
            else:
                user=authenticate(u,pw)
                if user:
                    clear_login_failures(u); st.session_state.authenticated=True; st.session_state.user=user
                    st.session_state.session_started=time.time()
                    saved_user_event=db_get("current_event_" + user.get("username",""))
                    if saved_user_event:
                        try: save_event_to_state(deserialize_event_obj(saved_user_event))
                        except Exception: pass
                    else:
                        save_event_to_state(blank_event())
                    st.session_state.current_event_id=db_get("current_event_id_"+user.get("username",""))
                    st.session_state.force_password_change=bool(user.get('must_change_password'))
                    st.rerun()
                else:
                    record_login_failure(u); st.error("Invalid username/email or password.")
    with register_tab:
        with st.form("register_form"):
            reml=st.text_input("Email Address")
            run=st.text_input("Username")
            rp1=st.text_input("Password",type="password")
            rp2=st.text_input("Confirm Password",type="password")
            reg=st.form_submit_button("Create Account",type="primary",use_container_width=True)
        if reg:
            if not valid_email(reml): st.error("Only NFSU institutional email addresses ending in @nfsu.ac.in are allowed.")
            elif len(run.strip())<3 or len(run.strip())>40: st.error("Username must contain 3–40 characters.")
            elif not run.replace('_','').replace('-','').isalnum(): st.error("Username may use letters, numbers, underscore and hyphen only.")
            elif not rp1 or rp1!=rp2: st.error("Password is required and both password fields must match.")
            else:
                ok,err=create_user(run,reml,rp1,'user')
                if ok: st.success("Account created. You can now sign in.")
                elif "email" in err.lower(): st.error("That email address is already registered.")
                else: st.error("That username is already registered.")
    st.markdown('</div>',unsafe_allow_html=True)
    st.markdown('<div class="site-footer"><strong>Developed by NFSU Goa Coding Club</strong><br>Coder: Dr. Ranjit Kolkar</div>',unsafe_allow_html=True)
    st.stop()

# Force first-login password change for accounts created/reset by an administrator.
if st.session_state.get('force_password_change'):
    st.markdown('<div class="page-title">Set Your Password</div><div class="page-subtitle">For account security, set a new password before continuing.</div>',unsafe_allow_html=True)
    with st.form('mandatory_password_change'):
        p1=st.text_input('New Password',type='password')
        p2=st.text_input('Confirm New Password',type='password')
        if st.form_submit_button('Save Password',type='primary'):
            if not p1 or p1!=p2: st.error('Password is required and both password fields must match.')
            else:
                change_password(st.session_state.user['username'],p1); st.session_state.force_password_change=False; st.session_state.user['must_change_password']=0; st.success('Password updated.'); st.rerun()
    st.markdown('<div class="site-footer"><strong>Developed by NFSU Goa Coding Club</strong><br>Coder: Dr. Ranjit Kolkar</div>',unsafe_allow_html=True)
    st.stop()

# ---------------- Sidebar ----------------
with st.sidebar:
    st.markdown('<div style="font-size:14px;font-weight:800;letter-spacing:.4px;margin-bottom:18px;">📋 EVENT MATERIAL GENERATOR</div>',unsafe_allow_html=True)
    nav=['My Events','Event Entry','Schedule','Experts & Guests','Budget','Checklist','Generate Materials'] + (['Admin'] if st.session_state.user.get('role')=='admin' else [])
    page=st.radio('Navigation',nav,index=nav.index(st.session_state.get('nav_page','My Events')) if st.session_state.get('nav_page','My Events') in nav else 0,label_visibility='collapsed'); st.session_state.nav_page=page
    st.divider(); st.caption(f"Signed in as: {st.session_state.user.get('username')}")
    if st.button('💾 Save All Changes',use_container_width=True): save_current_persistent(); st.success('Saved to local database.')
    if st.button('Logout',use_container_width=True): st.session_state.authenticated=False; st.session_state.user=None; st.rerun()


st.markdown('<div class="app-header"><div><div class="brand">EVENT MATERIAL GENERATOR</div><div class="brand-line"></div></div><div style="font-size:13px;color:#687789;">Institutional Event Planning & Documentation</div></div>',unsafe_allow_html=True)

# ---------------- My Events ----------------
if page=='My Events':
    st.markdown('<div class="page-title">My Events</div><div class="page-subtitle">View saved events or start a new event.</div>',unsafe_allow_html=True)
    hist=list_event_history(st.session_state.user.get('username'))
    if not hist:
        st.markdown('<div class="section"><div class="section-title">No events found</div><p>Create your first event or load the built-in sample data.</p></div>',unsafe_allow_html=True)
        a,b=st.columns(2)
        with a:
            if st.button('＋ Create New Event',type='primary',use_container_width=True):
                ev=blank_event(); save_event_to_state(ev); st.session_state.current_event_id=None; save_current_persistent(); st.session_state.nav_page='Event Entry'; st.rerun()
        with b:
            if st.button('Load Sample Event',use_container_width=True):
                ev=demo_event(); save_event_to_state(ev); eid=save_event_history(ev,st.session_state.user.get('username')); st.session_state.current_event_id=eid; save_current_persistent(); st.success('Sample event loaded.'); st.rerun()
    else:
        if st.button('＋ Create New Event',type='primary'):
            ev=blank_event(); save_event_to_state(ev); st.session_state.current_event_id=None; db_set('current_event_id_'+st.session_state.user.get('username',''),None); save_current_persistent(); st.session_state.nav_page='Event Entry'; st.rerun()
        st.markdown('### Saved Events')
        for r in hist:
            st.markdown(f'<div class="history-row"><div class="history-title">{r["title"]}</div><div class="history-meta">{r.get("start_date","")} to {r.get("end_date","")} · Last updated {r.get("updated_at","")}</div></div>',unsafe_allow_html=True)
            c1,c2,c3=st.columns([2,1,1])
            with c1:
                if st.button('Open Event',key=f'open_{r["id"]}',use_container_width=True):
                    rec=get_event_history(r['id']); save_event_to_state(deserialize_event_obj(rec['payload'])); st.session_state.current_event_id=r['id']; db_set("current_event_id_"+st.session_state.user.get("username",""), r['id']); st.session_state.nav_page='Event Entry'; st.rerun()
            with c2:
                if st.button('Duplicate',key=f'dup_{r["id"]}',use_container_width=True):
                    rec=get_event_history(r['id']); ev=deserialize_event_obj(rec['payload']); eid=save_event_history(ev,st.session_state.user.get('username')); save_event_to_state(ev); st.session_state.current_event_id=eid; db_set('current_event_id_'+st.session_state.user.get('username',''),eid); st.session_state.nav_page='Event Entry'; st.rerun()
            with c3:
                if st.button('Delete',key=f'del_{r["id"]}',use_container_width=True):
                    delete_event_history(r['id']);
                    if st.session_state.get('current_event_id')==r['id']: st.session_state.current_event_id=None; save_event_to_state(blank_event())
                    st.rerun()

# ---------------- Event Entry ----------------
if page=='Event Entry':

    st.markdown('<div class="page-title">Event / Proposal Details</div><div class="page-subtitle">Enter the event once. Optional fields can be included or excluded. The saved master data drives every generated material.</div>',unsafe_allow_html=True)
    e=st.session_state.event
    st.markdown('<div class="section"><div class="section-title">1. Basic Programme Information</div>',unsafe_allow_html=True)
    a,b=st.columns(2)
    with a:
        ptype=st.selectbox('Type of Proposed Programme',['Workshop','e-Workshop','Seminar','Conference','FDP','Training Programme','Hackathon','Other'],index=['Workshop','e-Workshop','Seminar','Conference','FDP','Training Programme','Hackathon','Other'].index(e.get('programme_type','Workshop')))
        title=st.text_input('Title of Proposed Programme *',e.get('title',''))
        faculty_opts=st.session_state.faculty
        selected=[x for x in e.get('coordinators',[]) if x in faculty_opts]
        coord_map={x['name']:x.get('designation','') for x in st.session_state.faculty_master}
        coords=st.multiselect('Coordinator(s) / Convener *',faculty_opts,default=selected,format_func=lambda n:f"{n} — {coord_map.get(n,'')}".rstrip(' — '))
        coord_extra=st.text_input('Additional Coordinator Name',value='')
        cocords=st.multiselect('Co-Coordinator(s) / Convener',faculty_opts,default=[x for x in e.get('co_coordinators',[]) if x in faculty_opts],format_func=lambda n:f"{n} — {coord_map.get(n,'')}".rstrip(' — '))
    with b:
        start=st.date_input('Start Date *',e.get('start_date',date.today())); end=st.date_input('End Date *',e.get('end_date',start)); mode=st.selectbox('Mode',['Offline','Online','Hybrid'],index=['Offline','Online','Hybrid'].index(e.get('mode','Offline'))); venue=st.text_input('Venue',e.get('venue','')); participants=st.number_input('Expected Participants',min_value=1,value=int(e.get('participants',40)))
        head=st.selectbox('Event Head',st.session_state.event_heads,index=st.session_state.event_heads.index(e.get('event_head')) if e.get('event_head') in st.session_state.event_heads else 0)
    nature=st.multiselect('Nature of Programme',['Theoretical','Practical','Workshop','Hands-on','Hybrid'],default=e.get('nature_list',['Workshop']))
    st.markdown('</div>',unsafe_allow_html=True)

    st.markdown('<div class="section"><div class="section-title">2. Institutional & Academic Information</div>',unsafe_allow_html=True)
    a,b=st.columns(2)
    with a:
        inst_name=st.text_input('Institution / University',st.session_state.institution.get('name','')); inst_dept=st.text_input('Department / School',st.session_state.institution.get('department','')); inst_addr=st.text_area('Institution Address',st.session_state.institution.get('address',''))
        c1,c2=st.columns([.28,4.72]);
        with c1: inc_inf=st.checkbox('Include',value=bool(e.get('infrastructure')),key='inc_inf',label_visibility='collapsed')
        with c2: infrastructure=st.text_area('Available Infrastructure / Laboratory Facilities',e.get('infrastructure',''),key='infrastructure')
        infrastructure=infrastructure if inc_inf else ''
    with b:
        audience=st.text_input('Targeted Audience',e.get('audience','')); lodging=st.text_area('Lodging Facilities',e.get('lodging','')); sponsors=st.text_area('Expected Sponsoring / Financing Agencies',e.get('sponsors','')); platform=st.text_input('Online Platform (if online / hybrid)',e.get('platform','')) if mode in ['Online','Hybrid'] else ''; other_info=st.text_area('Other Information',e.get('other_info',''))
    st.markdown('</div>',unsafe_allow_html=True)

    st.markdown('<div class="section"><div class="section-title">3. Compulsory Note Content</div><div class="small-muted">These paragraphs form the narrative master used by the Note, proposal and invitations.</div>',unsafe_allow_html=True)
    introduction=st.text_area('Introduction / Context *',e.get('introduction',''),height=120); justification=st.text_area('Justification / Need *',e.get('justification',''),height=120); objectives=st.text_area('Objectives *',e.get('objectives',''),height=120); outcome=st.text_area('Expected Outcome',e.get('outcome',''),height=120)
    st.markdown('</div>',unsafe_allow_html=True)

    st.markdown('<div class="section"><div class="section-title">4. Event Data File</div>',unsafe_allow_html=True)
    c1,c2,c3=st.columns(3)
    with c1:
        if st.button('💾 Save / Update Event',type='primary',use_container_width=True):
            coord_list=coords+([coord_extra.strip()] if coord_extra.strip() else [])
            missing=[]
            if not title.strip():missing.append('Event name')
            if not coord_list:missing.append('Coordinator name')
            if not justification.strip():missing.append('Justification')
            if not objectives.strip():missing.append('Objectives')
            if clean_schedule(st.session_state.schedule_df).empty:missing.append('Schedule')
            if start>end:missing.append('Valid date range')
            if missing:st.error('Please complete: '+', '.join(missing))
            else:
                st.session_state.institution.update({'name':inst_name,'department':inst_dept,'address':inst_addr})
                ev={**e,'institution':st.session_state.institution.copy(),'programme_type':ptype,'title':title,'coordinators':coord_list,'co_coordinators':cocords,'event_head':head,'start_date':start,'end_date':end,'mode':mode,'venue':venue,'participants':participants,'nature_list':nature,'audience':audience,'infrastructure':infrastructure,'organizer_expertise':e.get('organizer_expertise',''),'introduction':introduction,'justification':justification,'objectives':objectives,'outcome':outcome,'lodging':lodging,'sponsors':sponsors,'platform':platform,'other_info':other_info,'schedule_df':st.session_state.schedule_df.copy(),'inauguration_df':st.session_state.inauguration_df.copy(),'experts':st.session_state.experts.copy(),'budget':st.session_state.budget.copy(),'checklist':st.session_state.checklist.copy(),'chief_guest':st.session_state.chief_guest.copy(),'coordinator_signature':st.session_state.coordinator_signature,'invitation_text':st.session_state.invitation_text}
                st.session_state.event=ev; save_current_persistent(); st.session_state.current_event_id=save_event_history(ev, st.session_state.user.get('username'), st.session_state.get('current_event_id')); db_set('current_event_id_'+st.session_state.user.get('username',''),st.session_state.current_event_id); st.session_state.nav_page='Event Entry'; st.success('Event saved successfully.')
    with c2:
        st.download_button('⬇️ Download Current Event Data',serialize_event(st.session_state.event),file_name='event_data.json',mime='application/json',use_container_width=True)
    with c3:
        up=st.file_uploader('Upload Event JSON',type=['json'],label_visibility='collapsed')
        if up is not None:
            try:
                loaded=load_event_dict(up.getvalue()); save_event_to_state(loaded); save_current_persistent(); st.session_state.current_event_id=save_event_history(loaded, st.session_state.user.get('username')); db_set('current_event_id_'+st.session_state.user.get('username',''),st.session_state.current_event_id); st.session_state.nav_page='Event Entry'; st.success('Event data loaded and saved.'); st.rerun()
            except Exception as ex: st.error(f'Could not load JSON: {ex}')
    st.markdown('</div>',unsafe_allow_html=True)

# ---------------- Schedule ----------------
elif page=='Schedule':
    st.markdown('<div class="page-title">Programme Schedule</div><div class="page-subtitle">Build a clean schedule in a spreadsheet-style table. Column headers and rows are editable.</div>',unsafe_allow_html=True)
    a,b=st.columns([1,3]);
    with a: st.session_state.schedule_mode=st.selectbox('Schedule Format',['Day-wise','Combined'],index=['Day-wise','Combined'].index(st.session_state.schedule_mode))
    with b: st.info('Default columns: Day · Time · Topic · Expert Name. Use the table header menu to rename columns and the row controls to add/delete rows.',icon='ℹ️')
    st.markdown('<div class="section"><div class="section-title">Schedule Table</div>',unsafe_allow_html=True)
    edited=st.data_editor(st.session_state.schedule_df,use_container_width=True,num_rows='dynamic',hide_index=True,key='schedule_editor',height=420)
    st.session_state.schedule_df=edited
    save_current_persistent()
    c1,c2,c3=st.columns(3)
    with c1:
        new_col=st.text_input('New Column Header',placeholder='e.g. Venue',key='new_schedule_col')
        if st.button('＋ Add Column',use_container_width=True) and new_col.strip() and new_col.strip() not in st.session_state.schedule_df.columns:
            st.session_state.schedule_df[new_col.strip()]=''; st.rerun()
    with c2:
        if st.button('＋ Add 5 Rows',use_container_width=True):
            extra=pd.DataFrame([{c:'' for c in st.session_state.schedule_df.columns} for _ in range(5)]); st.session_state.schedule_df=pd.concat([st.session_state.schedule_df,extra],ignore_index=True); st.rerun()
    with c3:
        if st.button('↺ Reset Default',use_container_width=True): st.session_state.schedule_df=pd.DataFrame({'Day':['1','2','3','4','5'],'Time':['','','','',''],'Topic':['','','','',''],'Expert Name':['','','','','']}); st.rerun()
    st.markdown('</div>',unsafe_allow_html=True)
    st.markdown('<div class="section"><div class="section-title">Inauguration Programme</div>',unsafe_allow_html=True); st.session_state.inauguration_df=st.data_editor(st.session_state.inauguration_df,use_container_width=True,num_rows='dynamic',hide_index=True,key='inaug_editor',height=250); save_current_persistent(); st.markdown('</div>',unsafe_allow_html=True)

# ---------------- Experts ----------------
elif page=='Experts & Guests':
    st.markdown('<div class="page-title">Experts & Guests</div><div class="page-subtitle">Maintain probable experts and invitation details once. The proposal and invitations use this data.</div>',unsafe_allow_html=True)
    with st.form('expert_form',clear_on_submit=True):
        a,b=st.columns(2)
        with a:n=st.text_input('Expert Name'); aff=st.text_input('Affiliation'); spec=st.text_input('Field of Interest / Specialization')
        with b:topic=st.text_input('Proposed Topic'); role=st.text_input('Role / Session')
        if st.form_submit_button('＋ Add Expert',type='primary') and n.strip():st.session_state.experts.append({'name':n,'affiliation':aff,'specialization':spec,'topic':topic,'role':role});st.rerun()
    if st.session_state.experts:
        ed=st.data_editor(pd.DataFrame(st.session_state.experts),use_container_width=True,num_rows='dynamic',hide_index=True,key='expert_editor'); st.session_state.experts=ed.to_dict('records'); save_current_persistent()
    st.markdown('<div class="section"><div class="section-title">Chief Guest / Special Invitee</div>',unsafe_allow_html=True)
    g=st.session_state.chief_guest; a,b=st.columns(2)
    with a:gn=st.text_input('Name',g.get('name',''));gd=st.text_input('Designation',g.get('designation',''));ga=st.text_input('Affiliation',g.get('affiliation',''))
    with b:gaddr=st.text_area('Address',g.get('address',''));gr=st.text_input('Invitation Role',g.get('role','Chief Guest'))
    st.session_state.chief_guest={'name':gn,'designation':gd,'affiliation':ga,'address':gaddr,'role':gr};save_current_persistent();st.session_state.coordinator_signature=st.text_input('Coordinator Signature Block',st.session_state.coordinator_signature);st.session_state.invitation_text=st.text_area('Optional Invitation Paragraph',st.session_state.invitation_text);st.markdown('</div>',unsafe_allow_html=True)

# ---------------- Budget ----------------
elif page=='Budget':
    st.markdown('<div class="page-title">Budget Estimates</div><div class="page-subtitle">The total budget is automatically used as the amount requested in the generated Note and proposal.</div>',unsafe_allow_html=True)
    with st.form('budget_form',clear_on_submit=True):
        a,b,c=st.columns([4,2,4]);
        with a:item=st.text_input('Item Description')
        with b:amount=st.number_input('Amount (₹)',min_value=0.0,step=500.0)
        with c:remarks=st.text_input('Remarks')
        if st.form_submit_button('＋ Add Budget Item',type='primary') and item.strip():st.session_state.budget.append({'item':item,'amount':amount,'remarks':remarks});st.rerun()
    if st.session_state.budget:
        ed=st.data_editor(pd.DataFrame(st.session_state.budget),use_container_width=True,num_rows='dynamic',hide_index=True,key='budget_editor',column_config={'amount':st.column_config.NumberColumn('Amount (₹)',min_value=0,format='₹%.2f')});st.session_state.budget=ed.to_dict('records'); save_current_persistent(); total=sum(float(x.get('amount',0) or 0) for x in st.session_state.budget);st.metric('Total Expenses Requested for Approval',money(total));st.caption(f'In words: Rupees {amount_words(total)}')

# ---------------- Checklist ----------------
elif page=='Checklist':
    st.markdown('<div class="page-title">Event Checklist</div><div class="page-subtitle">Module-wise checklist for Pre-Event, In-Event and Post-Event. Each module is printable on A4 Word format.</div>',unsafe_allow_html=True)
    tabs=st.tabs(list(CHECKLISTS.keys())); statuses=['Not Started','In Progress','Completed','Delayed','Not Applicable']
    for tab,module in zip(tabs,CHECKLISTS):
        with tab:
            if module not in st.session_state.checklist:st.session_state.checklist[module]=[{'task':x,'status':'Not Started','responsible':'','due':'','remarks':''} for x in CHECKLISTS[module]]
            ed=st.data_editor(pd.DataFrame(st.session_state.checklist[module]),use_container_width=True,num_rows='dynamic',hide_index=True,key=f'check_{module}',column_config={'status':st.column_config.SelectboxColumn('Status',options=statuses,required=True)});st.session_state.checklist[module]=ed.to_dict('records'); save_current_persistent(); done=sum(x.get('status')=='Completed' for x in st.session_state.checklist[module]);total=len(st.session_state.checklist[module]);st.progress(done/total if total else 0);st.write(f'**{module}: {done}/{total} completed ({done/total*100 if total else 0:.0f}%)**');st.download_button(f'🖨️ Download {module} A4 Checklist',doc_bytes(build_checklist_doc(st.session_state.event,module)),file_name=f'{module.replace("-","_")}_Checklist.docx',mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',key=f'dl_{module}')

# ---------------- Generate ----------------
elif page=='Generate Materials':
    st.markdown('<div class="page-title">Generate Event Materials</div><div class="page-subtitle">All material is generated from the saved master event data. Checklist documents are included in the package.</div>',unsafe_allow_html=True)
    e=st.session_state.event; total=sum(float(x.get('amount',0) or 0) for x in st.session_state.budget); cards=st.columns(4)
    for c,label,val in zip(cards,['Event','Schedule Rows','Experts','Budget Requested'],[e.get('title',''),len(clean_schedule(st.session_state.schedule_df)),len(st.session_state.experts),money(total)]):
        with c:st.markdown(f'<div class="card"><div class="card-title">{label}</div><div class="card-value">{val}</div></div>',unsafe_allow_html=True)
    missing=[]
    if not e.get('title'):missing.append('Event name')
    if not e.get('coordinators'):missing.append('Coordinator')
    if not e.get('justification'):missing.append('Justification')
    if not e.get('objectives'):missing.append('Objectives')
    if clean_schedule(st.session_state.schedule_df).empty:missing.append('Schedule')
    if missing:st.error('Cannot generate yet. Missing: '+', '.join(missing))
    else:
        e={**e,'schedule_df':st.session_state.schedule_df.copy(),'inauguration_df':st.session_state.inauguration_df.copy(),'experts':st.session_state.experts.copy(),'budget':st.session_state.budget.copy(),'checklist':st.session_state.checklist.copy()}
        if st.button('📦 Generate Complete Event Package',type='primary',use_container_width=True):st.session_state.generated=build_package(e);st.success('Event package generated successfully.')
        if st.session_state.generated:
            z,files=st.session_state.generated;st.download_button('⬇️ Download Complete Event ZIP',z,file_name='Event_Material_Package.zip',mime='application/zip',use_container_width=True);st.subheader('Generated Materials')
            for name,data in files.items():st.download_button(f'⬇️ {name}',data,file_name=name,mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',key=f'gen_{name}')

# ---------------- Admin ----------------
elif page=='Admin':
    st.markdown('<div class="page-title">Administration / Master Data</div><div class="page-subtitle">Persistent master data and user management. Changes survive application restarts.</div>',unsafe_allow_html=True)
    st.markdown('<div class="section"><div class="section-title">Institutional Leadership</div>',unsafe_allow_html=True)
    ldf=pd.DataFrame([{'Role':k,'Name':v} for k,v in st.session_state.leadership.items()]); led=st.data_editor(ldf,use_container_width=True,num_rows='dynamic',hide_index=True,key='leadership_editor')
    st.session_state.leadership={str(r['Role']).strip():str(r['Name']).strip() for _,r in led.iterrows() if str(r.get('Role','')).strip() and str(r.get('Name','')).strip()}; save_current_persistent(); st.markdown('</div>',unsafe_allow_html=True)
    a,b=st.columns(2)
    with a:
        st.markdown('<div class="section"><div class="section-title">Faculty Master List</div>',unsafe_allow_html=True)
        fed=pd.DataFrame(st.session_state.faculty_master or [],columns=['name','designation']).rename(columns={'name':'Faculty Name','designation':'Designation'})
        fed=st.data_editor(fed,use_container_width=True,num_rows='dynamic',hide_index=True,key='faculty_editor')
        st.session_state.faculty_master=[{'name':str(r['Faculty Name']).strip(),'designation':str(r['Designation']).strip()} for _,r in fed.iterrows() if str(r.get('Faculty Name','')).strip()]
        st.session_state.faculty=[x['name'] for x in st.session_state.faculty_master]; save_current_persistent(); st.markdown('</div>',unsafe_allow_html=True)
    with b:
        st.markdown('<div class="section"><div class="section-title">Event Heads</div>',unsafe_allow_html=True)
        ed=st.data_editor(pd.DataFrame({'Event Head':st.session_state.event_heads}),use_container_width=True,num_rows='dynamic',hide_index=True,key='head_editor'); st.session_state.event_heads=[x for x in ed['Event Head'].fillna('').astype(str).tolist() if x.strip()]; save_current_persistent(); st.markdown('</div>',unsafe_allow_html=True)
    st.markdown('<div class="section"><div class="section-title">User Accounts</div>',unsafe_allow_html=True)
    users=db_users(); st.dataframe(pd.DataFrame(users),use_container_width=True,hide_index=True)
    with st.form('add_user_form'):
        x,y,z,w=st.columns(4)
        with x:nu=st.text_input('Username')
        with y:ne=st.text_input('Email')
        with z:np=st.text_input('Initial Password',type='password')
        with w:nr=st.selectbox('Role',['user','admin'])
        if st.form_submit_button('Create User',type='primary'):
            if not nu.strip() or not valid_email(ne) or not np: st.error('Username, valid NFSU email and password are required.')
            else:
                ok,err=create_user(nu,ne,np,nr,True)
                if ok: st.success('User created.'); st.rerun()
                elif 'email' in err.lower(): st.error('Email already exists.')
                else: st.error('Username already exists.')
    if users:
        st.markdown('#### Manage User')
        uid=st.selectbox('User Account',[u['id'] for u in users],format_func=lambda x: next((u['username'] for u in users if u['id']==x),str(x)))
        selected_user=next(u for u in users if u['id']==uid)
        a,b,c=st.columns(3)
        with a: ue=st.text_input('Email',selected_user.get('email') or '',key='manage_email')
        with b: ur=st.selectbox('Role',['user','admin'],index=['user','admin'].index(selected_user['role']),key='manage_role')
        with c: ua=st.checkbox('Active',value=bool(selected_user['active']),key='manage_active')
        q1,q2,q3=st.columns(3)
        with q1:
            if st.button('Save User',use_container_width=True):
                if not valid_email(ue): st.error('Enter a valid email address.')
                else: update_user(uid,ue,ur,ua); st.success('User updated.'); st.rerun()
        with q2:
            rp=st.text_input('Reset Password',type='password',key='reset_pw')
            if st.button('Reset Password',use_container_width=True):
                if not rp: st.error('Password is required.')
                else: reset_user_password(uid,rp); st.success('Password reset. User will set a new password at next sign-in.'); st.rerun()
        with q3:
            if st.button('Delete User',use_container_width=True):
                if selected_user['username']==st.session_state.user['username']: st.error('You cannot delete your own account.')
                else:
                    con=db_conn(); con.execute('DELETE FROM users WHERE id=?',(uid,)); con.commit(); con.close(); st.success('User deleted.'); st.rerun()
    st.markdown('<div class="section"><div class="section-title">Change My Password</div>',unsafe_allow_html=True)
    with st.form('change_password_form'):
        op=st.text_input('Current Password',type='password')
        p1=st.text_input('New Password',type='password')
        p2=st.text_input('Confirm New Password',type='password')
        if st.form_submit_button('Update Password'):
            con=db_conn(); row=con.execute('SELECT password_hash FROM users WHERE username=?',(st.session_state.user['username'],)).fetchone(); con.close()
            if not row or not verify_password(op,row['password_hash']): st.error('Current password is incorrect.')
            elif not p1 or p1!=p2: st.error('Password is required and both password fields must match.')
            else:
                con=db_conn(); con.execute('UPDATE users SET password_hash=? WHERE username=?',(hash_password(p1),st.session_state.user['username'])); con.commit(); con.close(); st.success('Password updated.')
    st.markdown('</div>',unsafe_allow_html=True)
    st.markdown('<div class="section"><div class="section-title">Event History</div>',unsafe_allow_html=True)
    hist=list_event_history()
    if hist:
        hdf=pd.DataFrame(hist)
        st.dataframe(hdf.rename(columns={'id':'ID','title':'Event','start_date':'Start Date','end_date':'End Date','owner_username':'Owner','created_at':'Created','updated_at':'Updated'}),use_container_width=True,hide_index=True)
        st.caption('Administrators can inspect and manage the complete stored event payload for every user.')
        ids=[int(x['id']) for x in hist]
        selected_id=st.selectbox('Select Event Record',ids,format_func=lambda x: next((f"{r['title']} — {r['start_date']}" for r in hist if r['id']==x),str(x)))
        selected_record=get_event_history(selected_id)
        with st.expander('View complete selected event data'):
            if selected_record: st.json(json.loads(selected_record['payload']))
        r1,r2,r3,r4=st.columns(4)
        with r1:
            if st.button('Load Event',use_container_width=True):
                rec=get_event_history(selected_id); ev=deserialize_event_obj(rec['payload']); save_event_to_state(ev); st.session_state.current_event_id=selected_id; db_set("current_event_id_"+st.session_state.user.get("username",""), selected_id); st.success('Event loaded.'); st.rerun()
        with r2:
            if st.button('Update Record',use_container_width=True):
                save_current_persistent(); st.session_state.current_event_id=selected_id; save_event_history(st.session_state.event,st.session_state.user.get('username'),selected_id); st.success('Event history record updated.'); st.rerun()
        with r3:
            if st.button('Duplicate as New',use_container_width=True):
                rec=get_event_history(selected_id); ev=deserialize_event_obj(rec['payload']); new_id=save_event_history(ev,st.session_state.user.get('username')); save_event_to_state(ev); st.session_state.current_event_id=new_id; db_set("current_event_id_"+st.session_state.user.get("username",""), new_id); st.success('Event duplicated.'); st.rerun()
        with r4:
            if st.button('Delete Record',use_container_width=True):
                delete_event_history(selected_id)
                if st.session_state.get('current_event_id')==selected_id: st.session_state.current_event_id=None
                st.success('Event history record deleted.'); st.rerun()
    else:
        st.info('No historical events have been saved yet.')
    st.markdown('</div>',unsafe_allow_html=True)




st.markdown('<div class="site-footer"><strong>Developed by NFSU Goa Coding Club</strong><br>Coder: Dr. Ranjit Kolkar</div>',unsafe_allow_html=True)
